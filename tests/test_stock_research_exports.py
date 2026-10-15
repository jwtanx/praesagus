"""Fictional export/lifecycle checks; no provider credentials, orders or fitting."""
import copy
import csv
from datetime import datetime,timezone,timedelta
import importlib
import importlib.util
import io
import json
from pathlib import Path
import subprocess
import sys

import pytest
from openpyxl import load_workbook

SCRIPT=Path(__file__).resolve().parents[1]/'scripts/stock-research'
sys.path.insert(0,str(SCRIPT))
# Other standalone test modules may replace storage in sys.modules. Reuse
# same-path modules, but load a new dependent module object for a stale class
# alias; never mutate older module objects retained by other test files.
for name in ('storage','collect','data','market_calendar','features','model','ranking','inference','evaluation','predict','export','demo','run'):
    existing=sys.modules.get(name)
    same=existing is not None and Path(getattr(existing,'__file__','')).resolve()==(SCRIPT/f'{name}.py').resolve()
    coherent=not hasattr(existing,'Generation') or existing.Generation is sys.modules['storage'].Generation
    if same and coherent:continue
    spec=importlib.util.spec_from_file_location(name,SCRIPT/f'{name}.py');module=importlib.util.module_from_spec(spec)
    sys.modules[name]=module;spec.loader.exec_module(module)
export=sys.modules['export'];demo=sys.modules['demo'];storage=sys.modules['storage'];run=sys.modules['run']
STAMP=datetime(2026,10,16,1,tzinfo=timezone.utc)


def rehash(payload):
    for row in payload['forecasts']:row['forecast_id']=storage.sha(storage.canonical({k:v for k,v in row.items() if k!='forecast_id'}))


def test_demo_readback_typed_values_all_views_and_immutable_collision(tmp_path):
    path=demo.run_demo(tmp_path,generated_at=STAMP);original={p.name:p.read_bytes() for p in path.iterdir()}
    manifest=json.loads((path/'manifest.json').read_text());assert manifest['is_demo'] is True
    assert storage.manifests(tmp_path,component='exports')==[]
    with pytest.raises(FileExistsError):demo.run_demo(tmp_path,generated_at=STAMP)
    assert {p.name:p.read_bytes() for p in path.iterdir()}==original
    universe,state,payload,evaluation=demo.fixtures();views,diagnostic,coverage=export.tables(universe,state,payload,evaluation,is_demo=True)
    for name,rows in views.items():
        actual=list(csv.DictReader(io.StringIO((path/(name+'.csv')).read_text())))
        assert len(actual)==len(rows)
        for a,b in zip(actual,rows):
            for key,value in b.items():assert a[key]==('' if value is None else str(export.literal(value)))
        assert all(r['is_demo']=='True' for r in actual)
    workbook=load_workbook(path/'research.xlsx')
    assert workbook.sheetnames[:5]==['Summary','Master','Winners','Losers','Volume']
    assert workbook.sheetnames[-6:]==['ETFs','Speculative','Coverage','Forecast review','Cumulative review','Forecasts']
    assert workbook['Summary'].max_column==8
    for i,(action,color) in enumerate(export.ACTIONS.items(),2):
        assert workbook['Summary'].cell(9,i).value==action
        assert workbook['Summary'].cell(9,i).fill.fgColor.rgb.endswith(color)
    assert '[' not in workbook['Summary']['H14'].value
    assert any(c.value=='NO TRADE' for row in workbook['Summary'] for c in row)
    assert any(c.value=='WATCH' for row in workbook['Summary'] for c in row)
    assert all(c.data_type!='f' for sheet in workbook for row in sheet for c in row)
    for sheet in workbook:
        assert sheet.freeze_panes=='D5' or sheet.title=='Summary'
        assert sheet.auto_filter.ref and sheet.sheet_view.showGridLines is False
    ws=workbook['Master'];headers=[c.value for c in ws[4]];rows=list(ws.iter_rows(min_row=5))
    assert headers[:5]==['code','name','asset_type','sector_id','action']
    assert rows[1][headers.index('name')].value=='=Literal name'
    assert rows[1][headers.index('name')].data_type=='s'
    eps=rows[1][headers.index('quote_financial_metrics_earning_per_share_value')]
    assert eps.value==-2.5 and eps.data_type=='n'
    assert all(row[headers.index('native_bullish_count')].value is None for row in rows)
    assert rows[0][headers.index('reference_session_date')].value.date().isoformat()=='2026-10-14'
    assert len(views['forecasts'])==20 and not views['top20'] and len(diagnostic)==4
    assert coverage['validated_candidates']==0
    assert all((path/name).stat().st_mode&0o777==0o600 for name in original)


def test_long_ids_all_flattened_workbook_identity_columns_are_literal_text():
    universe,state,payload,evaluation=demo.fixtures();item=universe['main'][0];code=item['code'];item['stock_id']=9223372036854775807
    state['identities'][code]['stock_id']=item['stock_id'];state['quotes'][code]['stock_id']=item['stock_id']
    for row in payload['forecasts']:
        if row['code']==code:row['stock_id']=item['stock_id']
    rehash(payload)
    views,diag,coverage=export.tables(universe,state,payload,evaluation,is_demo=True)
    w=load_workbook(io.BytesIO(export.workbook_bytes(views,diag,coverage)))
    sheet=w['Master'];headers=[c.value for c in sheet[4]]
    keys=[i for i,k in enumerate(headers) if k=='stock_id' or k.endswith('_stock_id')]
    assert len(keys)>4
    for i in keys:assert sheet.cell(5,i+1).value==str(item['stock_id']) and sheet.cell(5,i+1).data_type=='s'


@pytest.mark.parametrize('case',['missing-id','stale-id','string-id','identity-stale'])
def test_quote_financials_require_exact_typed_current_security(case):
    u,s,p,e=demo.fixtures();code=u['main'][0]['code']
    if case=='missing-id':s['quotes'][code].pop('stock_id')
    elif case=='stale-id':s['quotes'][code]['stock_id']+=1
    elif case=='string-id':s['quotes'][code]['stock_id']=str(s['quotes'][code]['stock_id'])
    else:s['identities'][code]['stock_id']+=1
    views,_,_=export.tables(u,s,p,e,is_demo=True);row=views['master'][0]
    assert row['quote_low_price'] is None and row['quote_high_price'] is None
    assert row['quote_status']!='bound-current-only'
    assert not any(k.startswith('quote_financial_metrics_') for k in row)


@pytest.mark.parametrize('case',['typed-id','row-id','duplicate','missing-window','window-date','window-time','started','origin',
    'axes','shape','price','terminal','crossing','nan','baseline','model-version','data-version','candidate-count'])
def test_invalid_forecast_contract_rejects_before_any_artifact(case,tmp_path):
    u,s,p,e=demo.fixtures();r=p['forecasts'][0]
    if case=='typed-id':r['instrument']=dict(r['instrument'],stock_id=str(r['instrument']['stock_id']))
    elif case=='row-id':r['stock_id']=str(r['stock_id'])
    elif case=='duplicate':p['forecasts'].append(copy.deepcopy(r))
    elif case=='missing-window':p['forecasts'].pop(0)
    elif case=='window-date':r['window']['end']='2026-10-19'
    elif case=='window-time':r['window']['first_open_ms']+=1
    elif case=='started':r['window']['started']=True
    elif case=='origin':r['reference']['session_date']='2026-10-13'
    elif case=='axes':r['model']=copy.deepcopy(r['model']);r['model']['targets'].reverse()
    elif case=='shape':r['model']=copy.deepcopy(r['model']);r['model']['price_quantiles'][0].pop()
    elif case=='price':r['model']=copy.deepcopy(r['model']);r['model']['price_quantiles'][0][0]=-1.
    elif case=='terminal':r['model']=copy.deepcopy(r['model']);r['model']['terminal_return_quantiles'][1]=.6
    elif case=='crossing':r['model']=copy.deepcopy(r['model']);r['model']['relative_quantiles'][0]=[.2,.1,.3]
    elif case=='nan':r['model']=copy.deepcopy(r['model']);r['model']['price_quantiles'][0][0]=float('nan')
    elif case=='baseline':r['baseline']=None
    elif case=='model-version':r['model_ref']={'different':'version'}
    elif case=='data-version':r['input_data_sha256']='different'
    elif case=='candidate-count':p['validated_candidate_count']=1
    if case!='nan':rehash(p)
    generation=storage.Generation(tmp_path,generated_at=STAMP,kind='stock-research-exports')
    with pytest.raises((ValueError,KeyError,TypeError,IndexError)):export.export_bundle(generation,u,s,p,e,is_demo=True)
    assert list(generation.path.iterdir())==[]


def test_all_unavailable_numeric_schema_remains_blank_without_zero_or_na():
    u,s,p,e=demo.fixtures()
    for row in p['forecasts']:row.update(status='unavailable',model=None,baseline=None,reference=None)
    rehash(p);views,_,_=export.tables(u,s,p,e,is_demo=True)
    for row in views['master']:
        for horizon in export.HORIZONS:
            assert row[horizon+'_low_median'] is None
            assert row[horizon+'_model_low_price_quantiles_q50'] is None
    raw=export.csv_bytes(views['forecasts'],export.columns(views['forecasts']))
    assert 'N/A' not in raw.decode()
    assert all(r['model_low_price_quantiles_q50']=='' for r in csv.DictReader(io.StringIO(raw.decode())))


def test_sector_names_collision_traversal_and_formula_csv_escape():
    names=['a/b','a:b','A:B','..','x'*100,'x'*99+'Y']
    tokens=[export.sector_token(n) for n in names]
    assert len({t.casefold() for t in tokens})==len(tokens) and all('/' not in t and len(t)<=31 for t in tokens)
    rows=[{'name':x,'negative':-2.5,'zero':0,'missing':None} for x in ['=1+2','+cmd','-name','@value',"'literal"]]
    parsed=list(csv.DictReader(io.StringIO(export.csv_bytes(rows,rows[0].keys()).decode())))
    assert all(r['name'].startswith("'") and r['negative']=='-2.5' and r['zero']=='0' and r['missing']=='' for r in parsed)


def test_export_collision_does_not_overwrite_or_finalize(tmp_path):
    g=storage.Generation(tmp_path,generated_at=STAMP,kind='stock-research-exports');g.write('master.csv',b'prior')
    u,s,p,e=demo.fixtures()
    with pytest.raises(ValueError,match='collision'):export.export_bundle(g,u,s,p,e,is_demo=True)
    assert (g.path/'master.csv').read_bytes()==b'prior' and not (g.path/'manifest.json').exists()


def test_run_empty_offline_one_collection_one_finish_and_no_provider(tmp_path,monkeypatch):
    settings=run.load_collection_config();from dataclasses import replace
    settings=replace(settings,state_parent=str(tmp_path));calls=[];finishes=[]
    original=run.collect_into;finish=storage.Generation.finish
    def collect_once(*args,**kwargs):calls.append(True);return original(*args,**kwargs)
    def finish_once(self,**kwargs):finishes.append(True);return finish(self,**kwargs)
    monkeypatch.setattr(run,'collect_into',collect_once);monkeypatch.setattr(storage.Generation,'finish',finish_once)
    cutoff=int(STAMP.timestamp()*1000)
    path=run.run_bundle(settings,offline=True,generated_at=STAMP,as_of_ms=cutoff,clock_ms=lambda:cutoff)
    m=json.loads((path/'manifest.json').read_text())
    assert calls==[True] and finishes==[True] and len(list(tmp_path.iterdir()))==1
    assert set(m['components'])=={'collection','prediction','exports'} and m['training'] is False
    assert (path/'history.csv').read_text().splitlines()[0]==','.join(storage.HISTORY_COLUMNS)
    assert (path/'predictions.json').exists() and (path/'research.xlsx').exists()
    assert storage.latest_state(tmp_path,run.compatibility(settings))['folder']==str(path)
    assert not (path/'universe.csv').exists()


def test_entrypoints_help_and_offline_are_credential_free(tmp_path):
    for script in ('run.py','demo.py'):
        result=subprocess.run([sys.executable,str(SCRIPT/script),'--help'],capture_output=True,text=True,timeout=15)
        assert result.returncode==0 and result.stderr==''
    result=subprocess.run([sys.executable,str(SCRIPT/'run.py'),'--offline','--sentiment','off','--state-dir',str(tmp_path)],capture_output=True,text=True,timeout=30)
    assert result.returncode==0,result.stdout+result.stderr
    path=Path(json.loads(result.stdout)['generation']);manifest=json.loads((path/'manifest.json').read_text())
    assert manifest['offline'] is True and manifest['source_capabilities']=={'live':'not-called'}
    assert manifest['sentiment_status']=='off'
    assert json.loads((path/'export_coverage.json').read_text())['sentiment_status']=='off'


def test_primary_compact_summary_preserves_supplied_action_and_eligibility():
    u,s,p,e=demo.fixtures();row=next(r for r in p['forecasts'] if r['horizon']=='3d' and r['code']==u['main'][0]['code'])
    row['decision'].update(action='BUY',score=1.234,roundtrip_cost=.002)
    row['validation']['qualified']=True;rehash(p);p['top20']=[row];p['validated_candidate_count']=1
    views,diag,coverage=export.tables(u,s,p,e,is_demo=True)
    assert views['top20'][0]['action']=='BUY' and views['top20'][0]['decision_score']==1.234
    w=load_workbook(io.BytesIO(export.workbook_bytes(views,diag,coverage)))
    assert w['Summary'].max_column==8 and w['Summary']['C8'].value=='BUY'
    assert w['Summary']['C8'].fill.fgColor.rgb.endswith(export.ACTIONS['BUY'])
    row['decision']['roundtrip_cost']=None;rehash(p)
    with pytest.raises(ValueError,match='primary-cost'):export.tables(u,s,p,e,is_demo=True)


def test_evaluation_metric_arrays_expose_numeric_columns():
    value={'metrics':{'model':{'low_return':{'pinball':[0.,.2,.3]}}}}
    result=export.flat(value)
    assert result['metrics_model_low_return_pinball_0']==0.
    assert result['metrics_model_low_return_pinball_1']==.2


@pytest.mark.parametrize('model_available',[True,False])
def test_nonempty_run_actual_inference_one_open_generation_and_no_fit(tmp_path,monkeypatch,model_available):
    # Reuse the accepted numerical fixture, not an authored output substituted for inference.
    import importlib.util
    path=SCRIPT.parents[1]/'tests/test_stock_research_inference.py'
    spec=importlib.util.spec_from_file_location('_prsg115_numerical_fixture',path)
    fixture=importlib.util.module_from_spec(spec);spec.loader.exec_module(fixture)
    calendar=fixture.MarketCalendar('2026-01-01','2027-04-01')
    history,instrument,contract,ref,cutoff,basis=fixture.fixture_inputs(calendar)
    actual_rows=[row for row in history.rows if row['known_at_ms']<=cutoff]
    settings=run.load_collection_config();from dataclasses import replace
    settings=replace(settings,state_parent=str(tmp_path));calls=[];predictions=[];finishes=[]
    receipt=history.receipts[0];state={'identities':{'US.DEMO':{'code':'US.DEMO','stock_id':123}},'quotes':{},
        'pools':{name:['US.DEMO'] for name in ('winners','losers','volume')}}
    universe={'main':[instrument],'speculative':[],'coverage':{}}
    from collect import CollectionResult
    def collected(*args,**kwargs):
        generation=kwargs['generation'];calls.append(generation)
        generation.csv('history.csv',actual_rows,storage.HISTORY_COLUMNS);generation.csv('history_revisions.csv',[],('code','session_date'))
        for name,value in [('identity_bindings.json',receipt['identity_bindings']),('state.json',{'schema_version':1,**state}),
            ('metadata_input.json',{}),('calendar.json',{}),('universe.json',universe),('coverage.json',{})]:generation.json(name,value)
        return CollectionResult(generation,copy.deepcopy(universe),state,actual_rows,receipt,{'as_of_ms':cutoff,'status':'complete','offline':True})
    predict=run.run_prediction;finish=storage.Generation.finish
    def predict_once(*args,**kwargs):predictions.append(kwargs['generation']);return predict(*args,**kwargs)
    def finish_once(self,**kwargs):finishes.append(self);return finish(self,**kwargs)
    monkeypatch.setattr(run,'collect_into',collected);monkeypatch.setattr(run,'run_prediction',predict_once)
    monkeypatch.setattr(storage.Generation,'finish',finish_once)
    model=fixture.ExistingModel() if model_available else None
    folder=run.run_bundle(settings,offline=True,model=model,contract=contract,model_ref=ref,
        basis_evidence=basis,costs={'US.DEMO':.001},as_of_ms=cutoff,clock_ms=lambda:cutoff,
        generated_at=datetime(2026,11,18,23,tzinfo=timezone.utc))
    assert len(calls)==len(predictions)==len(finishes)==1 and calls[0] is predictions[0] is finishes[0]
    assert len(list(tmp_path.iterdir()))==1
    payload=json.loads((folder/'predictions.json').read_text());assert len(payload['forecasts'])==4
    assert all(row['status']==('available' if model_available else 'unavailable') for row in payload['forecasts'])
    rows=list(csv.DictReader(io.StringIO((folder/'forecasts.csv').read_text())))
    assert len(rows)==4 and not list(csv.DictReader(io.StringIO((folder/'top20.csv').read_text())))
    if model_available:
        assert len(model.calls)==4 and all(float(row['model_low_price_quantiles_q50'])>0 for row in rows)
        assert storage.forecast_symbols(tmp_path)=={'US.DEMO'}
    else:assert all(row['model_low_price_quantiles_q50']=='' for row in rows)
    workbook=load_workbook(folder/'research.xlsx');sheet=workbook['Forecasts'];headers=[c.value for c in sheet[4]]
    for i,row in enumerate(rows,5):
        cell=sheet.cell(i,headers.index('model_low_price_quantiles_q50')+1)
        assert cell.value==pytest.approx(float(row['model_low_price_quantiles_q50'])) if model_available else cell.value is None
