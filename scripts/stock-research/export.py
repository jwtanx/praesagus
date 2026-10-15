"""Typed CSV/XLSX views only. No forecasting, fitting, providers or action policy."""
import csv
from datetime import date, timedelta
import hashlib
import io
import json
import math
import re

from storage import canonical, sha

HORIZONS=('1d','3d','5d','next_week')
ACTIONS={'BUY':'D9EAD3','SELL':'F4CCCC','WATCH':'FCE5CD','NO TRADE':'E7E6E6'}


def require(value,reason):
    if not value:raise ValueError(reason)


def identity(row):
    value=row.get('stock_id')
    require(type(value) is int and 0<value<=2**63-1 or isinstance(value,str) and re.fullmatch('[0-9]{1,19}',value) is not None and 0<int(value)<=2**63-1,'export-security-id-required')
    require(isinstance(row.get('code'),str) and re.fullmatch(r'US\.[A-Z][A-Z0-9.-]{0,63}',row['code']) is not None,'export-code')
    return row['code'],type(value).__name__,value


def flat(value,prefix=''):
    result={}
    for key,item in value.items():
        name=prefix+key
        if isinstance(item,dict):result.update(flat(item,name+'_'))
        elif isinstance(item,(list,tuple)):
            result[name]=json.dumps(item,sort_keys=True,ensure_ascii=False,allow_nan=False)
            if all(type(v) in (int,float) and math.isfinite(v) for v in item):
                result.update({name+'_'+str(i):v for i,v in enumerate(item)})
        else:
            require(type(item) is not float or math.isfinite(item),'export-nonfinite-value')
            result[name]=item
    return result


def literal(value):
    # CSV text escaping is reversible: every leading apostrophe is escaped too.
    return "'"+value if isinstance(value,str) and value[:1] in ('=','+','-','@',"'",'\t','\r','\n') else value


def csv_bytes(rows,columns):
    stream=io.StringIO(newline='');writer=csv.DictWriter(stream,fieldnames=columns,lineterminator='\n')
    writer.writeheader()
    for row in rows:writer.writerow({k:literal(row.get(k)) for k in columns})
    return stream.getvalue().encode()


def sector_token(value):
    text=str(value);stem=re.sub('[^A-Za-z0-9_-]+','_',text).strip('_')[:12] or 'unknown'
    return 'sector_'+stem+'_'+hashlib.sha256(text.encode()).hexdigest()[:10]


def forecast_row(row,is_demo):
    result={'code':row['code'],'stock_id':row['instrument'].get('stock_id'),
        'stock_id_type':type(row['instrument'].get('stock_id')).__name__,'name':row['instrument'].get('name'),
        'horizon':row['horizon'],'forecast_id':row['forecast_id'],'action':row['decision']['action'],
        'status':row['status'],'reference_close':None,'reference_session_date':None,'reference_currency':None,'reference_adjustment':None,'is_demo':is_demo}
    for family in ('model','baseline'):
        for field in ('relative_quantiles','price_quantiles'):
            for axis in ('low','high','terminal'):
                for q in ('q05','q50','q95'):result[f'{family}_{axis}_{field}_{q}']=None
        for q in ('q05','q50','q95'):result[f'{family}_terminal_return_{q}']=None
    result.update(flat({key:value for key,value in row.items() if key not in {'code','stock_id','horizon','forecast_id','is_demo'}}))
    for family in ('model','baseline'):
        values=row.get(family)
        if values is None:continue
        require(values.get('targets')==['low_return','high_return','terminal_logreturn'] and values.get('quantiles')==[.05,.5,.95],'export-quantile-contract')
        for field in ('relative_quantiles','price_quantiles'):
            matrix=values[field]
            require(len(matrix)==3 and all(len(axis)==3 for axis in matrix),'export-quantile-shape')
            for axis,name in enumerate(('low','high','terminal')):
                for q,label in enumerate(('q05','q50','q95')):
                    value=matrix[axis][q]
                    require(type(value) in (int,float) and math.isfinite(value),'export-invalid-quantile')
                    result[f'{family}_{name}_{field}_{label}']=value
        relative=values['relative_quantiles'];prices=values['price_quantiles'];returns=values['terminal_return_quantiles']
        require(len(returns)==3,'export-terminal-shape')
        require(row.get('reference') and type(row['reference'].get('close')) in (float,int) and row['reference']['close']>0,'export-reference-price')
        reference=row['reference']['close']
        for axis in range(3):
            require(all(relative[axis][q]<=relative[axis][q+1] for q in range(2)),'export-crossing-quantiles')
            for q in range(3):
                try:expected=reference*(1+relative[axis][q]) if axis<2 else reference*math.exp(relative[axis][q])
                except OverflowError:raise ValueError('export-terminal-overflow') from None
                require(math.isfinite(expected) and expected>0 and math.isclose(expected,prices[axis][q],rel_tol=1e-10,abs_tol=1e-10),'export-price-conversion-mismatch')
        require(relative[0][1]<=relative[1][1] and values.get('predicted_low')==prices[0][1] and values.get('predicted_high')==prices[1][1],'export-extrema-mismatch')
        for q,label in enumerate(('q05','q50','q95')):
            require(type(returns[q]) in (int,float) and math.isfinite(returns[q]) and math.isclose(returns[q],math.expm1(relative[2][q]),rel_tol=1e-10,abs_tol=1e-10),'export-terminal-conversion-mismatch')
            result[f'{family}_terminal_return_{label}']=returns[q]
    return result


def tables(universe,state,payload,evaluation,*,is_demo=False):
    require(payload.get('schema_version')==1 and payload.get('kind')=='stock-research-forecast-bundle','export-payload-schema')
    require(evaluation.get('schema_version')==1 and evaluation.get('kind')=='stock-research-evaluation','export-evaluation-schema')
    require(type(is_demo) is bool,'export-demo-flag')
    require(type(payload.get('as_of_ms')) is int and type(payload.get('issued_at_ms')) is int and 0<payload['as_of_ms']<=payload['issued_at_ms'],'export-cutoff-order')
    instruments=universe['main']+universe['speculative']
    require(len(instruments)<=1000 and len(payload['forecasts'])<=4000,'export-row-budget')
    by_code={r['code']:r for r in instruments}
    require(len(by_code)==len(instruments),'export-duplicate-instrument')
    require(len({(identity(r)[1],identity(r)[2]) for r in instruments})==len(instruments),'export-duplicate-security')
    require(payload.get('is_demo',is_demo)==is_demo and evaluation.get('is_demo',is_demo)==is_demo,'export-demo-mismatch')
    require(all(bool(r.get('is_demo',False))==is_demo for r in instruments+payload['forecasts']),'export-demo-mismatch')
    calendar=None
    if payload['forecasts']:
        from market_calendar import MarketCalendar
        first=min(date.fromisoformat(r['origin']) for r in payload['forecasts'])
        last=max(date.fromisoformat(r['origin']) for r in payload['forecasts'])
        calendar=MarketCalendar((first-timedelta(days=10)).isoformat(),(last+timedelta(days=40)).isoformat())
    seen=set();ids={};long=[];grouped={} 
    for row in payload['forecasts']:
        code=row['code'];require(code in by_code,'export-unknown-instrument')
        require(identity(row['instrument'])==identity(by_code[code]),'export-security-binding-mismatch')
        if row.get('stock_id') is not None:require(type(row['stock_id']) is type(by_code[code]['stock_id']) and row['stock_id']==by_code[code]['stock_id'],'export-forecast-security-mismatch')
        require(row['horizon'] in HORIZONS and (code,row['horizon']) not in seen,'export-duplicate-or-unknown-horizon')
        require(row['issued_at_ms']==payload['issued_at_ms'] and row['as_of_ms']==payload['as_of_ms'] and row['input_data_sha256']==payload['data_sha256'] and row['model_ref']==payload['model_ref'],'export-issuance-version-mismatch')
        require(row['forecast_id']==sha(canonical({k:v for k,v in row.items() if k!='forecast_id'})),'export-forecast-id-mismatch')
        require(row['decision']['action'] in ACTIONS,'export-action')
        require(row.get('status') in {'available','unavailable','invalid'},'export-forecast-status')
        if row['status']=='available':
            require(row.get('stock_id') is not None,'export-available-security-binding-required')
            require(row.get('model') is not None and row.get('baseline') is not None,'export-matched-baseline-required')
        window=row['window'];expected=calendar.windows(row['origin'])[row['horizon']]
        require(window['horizon']==row['horizon'] and window['sessions']==expected and window['start']==expected[0] and window['end']==expected[-1], 'export-window-mismatch')
        require(window.get('first_open_ms')==calendar.open_ms(expected[0]) and window.get('last_close_ms')==calendar.close_ms(expected[-1]) and type(window.get('started')) is bool and window['started']==(row['as_of_ms']>=calendar.open_ms(expected[0])),'export-window-times')
        require(calendar.close_ms(row['origin'])<=row['as_of_ms'],'export-reference-session-future')
        if row.get('reference') is not None:
            ref=row['reference']
            require(ref['session_date']==row['origin'] and type(ref.get('known_at_ms')) is int and calendar.close_ms(row['origin'])<=ref['known_at_ms']<=row['as_of_ms'],'export-reference-availability')
        seen.add((code,row['horizon']));ids[row['forecast_id']]=row
        item=forecast_row(row,is_demo);long.append(item);grouped.setdefault(code,{})[row['horizon']]=item
    for code in by_code:
        require(set(grouped.get(code,{}))==set(HORIZONS),'export-four-windows-required')
        references=[ids[item['forecast_id']]['reference'] for item in grouped[code].values()]
        require(all(r==references[0] for r in references),'export-reference-mismatch')
    master=[];speculative=[]
    for instrument in instruments:
        code=instrument['code'];row={'code':code,'stock_id':instrument['stock_id'],'stock_id_type':type(instrument['stock_id']).__name__,
            'name':instrument.get('name'),'asset_type':instrument.get('asset_type'),'sector_id':instrument.get('sector_id'),
            'signed_leverage':instrument.get('signed_leverage'),'issuer_group_id':instrument.get('issuer_group_id'),'is_demo':is_demo,
            'native_bullish_count':None,'native_bearish_count':None,'native_count_status':'methodology-unverified',
            'native_count_source':None,'native_count_timeframe':None}
        row.update(flat(instrument,'instrument_'))
        current=state.get('identities',{}).get(code)
        current_ok=isinstance(current,dict) and current.get('code')==code and type(current.get('stock_id')) is type(instrument['stock_id']) and current.get('stock_id')==instrument['stock_id']
        row['current_identity_status']='bound' if current_ok else 'unavailable-or-security-mismatch'
        if current_ok:row.update(flat({k:v for k,v in current.items() if k in {'code','stock_id','name','stock_type','exchange','listing_date','suspension','state','source','retrieved_at_ms','known_at_ms'}},'identity_'))
        quote=state.get('quotes',{}).get(code)
        for key in ('last_price','low_price','high_price','open_price','prev_close_price','volume','turnover','source_at_ms','known_at_ms','retrieved_at_ms'):row['quote_'+key]=None
        row['quote_status']='quote-security-id-unavailable'
        if isinstance(quote,dict) and quote.get('stock_id') is not None:
            quote_ok=current_ok and quote.get('code')==code and type(quote['stock_id']) is type(instrument['stock_id']) and quote['stock_id']==instrument['stock_id']
            row['quote_status']='bound-current-only' if quote_ok else 'quote-security-binding-mismatch'
            if quote_ok:row.update(flat({k:v for k,v in quote.items() if k in {'code','stock_id','source','data_date','session_date','last_price','low_price','high_price','open_price','prev_close_price','volume','turnover','source_at_ms','known_at_ms','retrieved_at_ms','financial_metrics'}},'quote_'))
        for horizon,item in grouped[code].items():row.update({horizon+'_'+key:value for key,value in item.items() if key not in {'code','stock_id','stock_id_type','name','is_demo'}})
        row['action']=grouped[code]['3d']['action'];row['status']=grouped[code]['3d']['status']
        reference=ids[grouped[code]['3d']['forecast_id']].get('reference') or {}
        front={'code':code,'name':instrument.get('name'),'asset_type':instrument.get('asset_type'),
            'sector_id':instrument.get('sector_id'),'action':row['action'],'reference_close':reference.get('close'),
            'reference_session_date':reference.get('session_date'),'quote_low_price':row['quote_low_price'],
            'quote_high_price':row['quote_high_price'],'quote_status':row['quote_status']}
        for horizon in HORIZONS:
            front[horizon+'_low_median']=grouped[code][horizon]['model_low_price_quantiles_q50']
            front[horizon+'_high_median']=grouped[code][horizon]['model_high_price_quantiles_q50']
        row=front|row
        (speculative if code in {m['code'] for m in universe['speculative']} else master).append(row)
    primary=[]
    for selected in payload.get('top20',[]):
        row=ids.get(selected['forecast_id']);require(row==selected,'export-rank-row-mismatch')
        require(row['horizon']=='3d' and row['status']=='available' and row['validation']['qualified'] is True and not row['window']['started'] and row['decision']['action'] in {'BUY','SELL'},'export-ineligible-primary')
        require(row['code'] in {m['code'] for m in universe['main']},'export-speculative-primary')
        cost=row['decision'].get('roundtrip_cost');score=row['decision'].get('score')
        require(type(cost) in (int,float) and math.isfinite(cost) and cost>=0 and type(score) in (int,float) and math.isfinite(score),'export-primary-cost-or-score')
        primary.append(grouped[row['code']]['3d'])
    require(len(primary)<=20 and len({r['code'] for r in primary})==len(primary),'export-top20-cap')
    require(payload.get('validated_candidate_count')==len(primary),'export-candidate-count')
    diagnostic=[grouped[r['code']]['3d'] for r in payload['forecasts'] if r['horizon']=='3d' and r['status']=='available' and r['decision']['action'] in {'WATCH','NO TRADE'} and r['code'] in {m['code'] for m in universe['main']}]
    result={'master':master,'forecasts':long,'top20':primary,'speculative':speculative,
        'etfs':[r for r in master if r.get('asset_type')=='ETF']}
    pools=state.get('pools',{})
    for name in ('winners','losers','volume'):
        rank={code:i for i,code in enumerate(pools.get(name,[]))}
        result[name]=sorted([r for r in master if r['code'] in rank],key=lambda r:rank[r['code']])
    for sector in sorted({str(r.get('sector_id') or 'unknown') for r in master}):result[sector_token(sector)]=[r for r in master if str(r.get('sector_id') or 'unknown')==sector]
    result['evaluation']=[flat(row)|{'is_demo':is_demo} for row in evaluation.get('outcome_revisions',[])]
    result['evaluation_cumulative']=[{'group':name,**flat(value),'is_demo':is_demo} for name,value in evaluation.get('cumulative',{}).get('groups',{}).items()]
    coverage={'schema_version':1,'is_demo':is_demo,'main_instruments':len(master),'speculative_instruments':len(speculative),
        'available_numerical_forecasts':sum(r['status']=='available' for r in payload['forecasts']),
        'validated_candidates':len(primary),'research_diagnostic_rows':len(diagnostic),'quality_status':payload.get('quality_status'),
        'collection':universe.get('coverage',{}),'sentiment_status':universe.get('coverage',{}).get('sentiment_status','unavailable/not-configured-until-PRSG-116')}
    return result,diagnostic,coverage


def columns(rows):
    keys=list(dict.fromkeys(key for row in rows for key in row))
    return keys or ['code','status','reason','is_demo']


def workbook_bytes(views,diagnostic,coverage):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.formatting.rule import CellIsRule
    workbook=Workbook();workbook.remove(workbook.active)
    def sheet(name,rows,start=4,cols_override=None):
        ws=workbook.create_sheet(name);ws.sheet_view.showGridLines=False
        ws.cell(2,1,name).font=Font(name='Arial',size=14,bold=True,color='17365D')
        cols=cols_override or columns(rows)
        for index,key in enumerate(cols,1):
            cell=ws.cell(start,index,key);cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            ws.column_dimensions[get_column_letter(index)].width=min(42,max(16,len(key)*.8))
        ws.row_dimensions[start].height=42
        for number,row in enumerate(rows,start+1):
            ws.row_dimensions[number].height=21
            for index,key in enumerate(cols,1):
                value=row.get(key)
                if (key=='stock_id' or key.endswith('_stock_id')) and value is not None:value=str(value)
                if isinstance(value,str) and (key.endswith(('_date','_start','_end','_origin')) or key=='origin'):
                    try:value=date.fromisoformat(value)
                    except ValueError:pass
                cell=ws.cell(number,index,value);cell.font=Font(name='Arial',size=10,color='243746');cell.alignment=Alignment(vertical='center',horizontal='right' if type(value) in (int,float) else 'left')
                if isinstance(value,str):cell.data_type='s'
                if type(value) is float:cell.number_format='0.0000' if any(word in key for word in ('return','ratio','cost','threshold')) else '#,##0.00'
                if type(value) is int:cell.number_format='#,##0'
                if isinstance(value,date):cell.number_format='yyyy-mm-dd'
                if isinstance(value,str) and value in ACTIONS:cell.fill=PatternFill('solid',fgColor=ACTIONS[value])
        ws.auto_filter.ref=f'A{start}:{get_column_letter(len(cols))}{max(start,len(rows)+start)}'
        ws.freeze_panes=f'D{start+1}'
        for index,key in enumerate(cols,1):
            if key=='action' or key.endswith('_action'):
                letter=get_column_letter(index)
                for action,color in ACTIONS.items():ws.conditional_formatting.add(f'{letter}{start+1}:{letter}{max(start+1,start+len(rows))}',CellIsRule(operator='equal',formula=['"'+action+'"'],fill=PatternFill('solid',fgColor=color)))
        return ws
    def compact(row):
        parts=[]
        for raw in (row.get('decision_reasons'),row.get('validation_reasons')):
            if raw:
                values=json.loads(raw) if isinstance(raw,str) and raw.startswith('[') else [raw]
                parts.extend(re.sub(r'(?<=[A-Za-z])-(?=[A-Za-z])',' ',str(value)).replace('_',' ') for value in values)
        reasons='; '.join(dict.fromkeys(parts))
        return {'Symbol':row['code'],'Name':row.get('name'),'Action':row['action'],
            'Reference close':row.get('reference_close'),'3d low median':row.get('model_low_price_quantiles_q50'),
            '3d high median':row.get('model_high_price_quantiles_q50'),'Score':row.get('decision_score'),
            'Validation / reasons':reasons}
    summary_columns=['Symbol','Name','Action','Reference close','3d low median','3d high median','Score','Validation / reasons']
    summary=sheet('Summary',[compact(row) for row in views['top20']],start=7,cols_override=summary_columns)
    for letter,width in zip('ABCDEFGH',[21,27,15,17,17,17,15,52]):summary.column_dimensions[letter].width=width
    summary.merge_cells('A3:H3');summary.merge_cells('A4:H4');summary.merge_cells('A5:H5')
    summary.cell(3,1,'Fictional template' if coverage['is_demo'] else 'Market research')
    summary.cell(4,1,f"Validated candidates: {coverage['validated_candidates']}   Numerical forecasts (all windows): {coverage['available_numerical_forecasts']}")
    summary.cell(5,1,coverage['quality_status'] or 'Validation unavailable').alignment=Alignment(wrap_text=True)
    summary.row_dimensions[5].height=32
    if not views['top20']:summary.cell(8,1,'No validated candidates. Numerical research is shown below.')
    legend=9+len(views['top20']);summary.cell(legend,1,'Action colors').font=Font(name='Arial',size=10)
    for index,(action,color) in enumerate(ACTIONS.items(),2):
        cell=summary.cell(legend,index,action);cell.fill=PatternFill('solid',fgColor=color);cell.font=Font(name='Arial',size=10)
    at=12+len(views['top20']);summary.merge_cells(start_row=at,start_column=1,end_row=at,end_column=8)
    summary.cell(at,1,'Research diagnostics: WATCH and NO TRADE; not validated candidates').font=Font(name='Arial',size=10,bold=True)
    summary.row_dimensions[at+1].height=32
    for i,key in enumerate(summary_columns,1):
        cell=summary.cell(at+1,i,key);cell.font=Font(name='Arial',bold=True,size=10,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='17365D');cell.alignment=Alignment(wrap_text=True,vertical='center')
    for n,row in enumerate(diagnostic,at+2):
        summary.row_dimensions[n].height=56
        for i,key in enumerate(summary_columns,1):
            cell=summary.cell(n,i,compact(row).get(key));cell.font=Font(name='Arial',size=10)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
            if isinstance(cell.value,str):cell.data_type='s'
            if type(cell.value) is float:cell.number_format='0.00'
            if key=='Action':cell.fill=PatternFill('solid',fgColor=ACTIONS[row['action']])
    order=[('Master','master'),('Winners','winners'),('Losers','losers'),('Volume','volume')]
    order += [(name[:31],name) for name in sorted(views) if name.startswith('sector_')]
    order += [('ETFs','etfs'),('Speculative','speculative'),('Coverage',None),('Forecast review','evaluation'),('Cumulative review','evaluation_cumulative'),('Forecasts','forecasts')]
    for title,key in order:sheet(title,views[key] if key else [flat(coverage)])
    out=io.BytesIO();workbook.save(out);return out.getvalue()


def export_bundle(generation,universe,state,payload,evaluation,*,is_demo=False):
    require(not generation.closed,'generation-already-finalized')
    views,diagnostic,coverage=tables(universe,state,payload,evaluation,is_demo=is_demo)
    names={key+'.csv' for key in views}|{'export_coverage.json','research.xlsx'}
    require(not names.intersection(generation.manifest['artifacts']) and not any((generation.path/name).exists() for name in names),'export-artifact-collision')
    binary=workbook_bytes(views,diagnostic,coverage)
    for key,rows in views.items():generation.write(key+'.csv',csv_bytes(rows,columns(rows)))
    generation.json('export_coverage.json',coverage);generation.write('research.xlsx',binary)
    return {'artifacts':sorted(names),'coverage':coverage}
